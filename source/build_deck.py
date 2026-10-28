# LoanShield deck: the team's Part 1 slides restyled + the solution slides, all in the title-slide style.
# python3 build_deck.py <team Part 1 .pptx> <out.pptx>
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

SRC, OUT = sys.argv[1], sys.argv[2]
IMG = '/home/claude/loanshield/deckimg/'
LINK = 'https://karthik705.github.io/loanshield/'
C = lambda h: RGBColor.from_string(h)
BG = '0E0807'
INK, SOFT, MUTE = 'F4EEEC', 'CFC4C3', '9C8E8C'
CARD, EDGE, LINE = '1A1313', '3A2A28', '2A1F1E'
RED, RED_DK, RED_BG = 'E63946', '5A2630', '2A0E10'
ORG, ORG_BG = 'F4A261', '2A1A10'
GLOW = '1F0A0B'
SDG = {1: 'E5243B', 8: 'A21942', 10: 'DD1367', 16: '00689D'}
F = 'Poppins'
L, R = 0.9, 12.43  # content left / right edge (13.33 wide)
CW = R - L
CENTER, MIDDLE = PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE

prs = Presentation(SRC)
BLANK = prs.slide_layouts[6]
page = [0]

def text(s, x, y, w, h, runs, size=13, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, gap=5):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = ((p, {}),) if isinstance(p, str) else p
        for t, o in segs:
            r = para.add_run(); r.text = t; f = r.font
            f.name = F; f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold); f.color.rgb = C(o.get('color', color))
        if i: para.space_before = Pt(gap)
        if spacing: para.line_spacing = spacing
    return tb

def shape(s, kind, x, y, w, h, fill, edge=None, lw=1.0, adj=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    if edge: sh.line.color.rgb = C(edge); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if adj is not None: sh.adjustments[0] = adj
    return sh

def box(s, x, y, w, h, fill=CARD, edge=EDGE, lw=1.0, adj=0.08):
    return shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, edge, lw, adj)

def oval(s, x, y, d, fill, edge=None): return shape(s, MSO_SHAPE.OVAL, x, y, d, d, fill, edge)
def rule(s, x, y, w, color=RED, h=0.06): return shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, color)

def pill(s, x, y, t, color=RED, fill=RED_BG, size=10):
    w = len(t) * 0.088 * size / 10 + 0.55
    box(s, x, y, w, 0.36, fill, color, 1, 0.5)
    text(s, x, y, w, 0.36, t, size=size, color=color, bold=True, align=CENTER, anchor=MIDDLE)
    return w

def tile(s, x, y, n, color, d=0.6, fs=17):
    box(s, x, y, d, d, color, None, adj=0.18)
    text(s, x, y, d, d, str(n), size=fs, bold=True, align=CENTER, anchor=MIDDLE)

def dot(s, x, y, n, color=RED, d=0.46, fs=12):
    oval(s, x, y, d, color)
    text(s, x, y, d, d, str(n), size=fs, bold=True, align=CENTER, anchor=MIDDLE)

def notes(s, t):
    if t: s.notes_slide.notes_text_frame.text = t

def glow(s, where):
    spots = {'tr': (8.6, -1.6, 6.2), 'br': (9.3, 3.6, 5.6), 'bl': (-2.4, 3.9, 5.6), 'tl': (-2.6, -2.4, 5.8), 'r': (9.6, 0.8, 6.0)}
    x, y, d = spots[where]; oval(s, x, y, d, GLOW)

def frame(eyebrow=None, title=None, where='tr', fresh=None):
    """Background, glow, eyebrow pill, title and footer for one slide."""
    if fresh is None: s = prs.slides.add_slide(BLANK)
    else:
        s = fresh
        for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = C(BG)
    glow(s, where)
    if eyebrow: pill(s, L, 0.55, eyebrow)
    if title: text(s, L, 1.05, CW, 0.9, title, size=30, bold=True)
    page[0] += 1
    rule(s, L, 7.0, CW, LINE, 0.01)
    text(s, L, 7.07, 6, 0.25, 'LoanShield  ·  Social & Sustainability', size=8.5, color=MUTE)
    text(s, R - 1, 7.07, 1, 0.25, str(page[0]).zfill(2), size=8.5, color=MUTE, align=PP_ALIGN.RIGHT)
    return s

def card(s, x, y, w, h, head, body, hc=RED, edge=EDGE, fill=CARD, hs=14, bs=11, num=None, ncol=RED):
    box(s, x, y, w, h, fill, edge)
    tx = x + 0.25
    if num is not None:
        dot(s, tx, y + 0.25, num, ncol, 0.42, 11)
        text(s, tx + 0.58, y + 0.25, w - 0.95, 0.42, head, size=hs, color=hc, bold=True, anchor=MIDDLE)
        text(s, tx, y + 0.82, w - 0.5, h - 0.95, body, size=bs, color=SOFT)
    else:
        text(s, tx, y + 0.24, w - 0.5, 0.5, head, size=hs, color=hc, bold=True)
        text(s, tx, y + 0.24 + hs / 72 * 1.55, w - 0.5, h - 0.7, body, size=bs, color=SOFT)

def picture(s, path, x, y, w=None, h=None):
    iw, ih = Image.open(path).size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    box(s, x - 0.08, y - 0.08, w + 0.16, h + 0.16, CARD, EDGE, 1, 0.04)
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h

def banner(s, y, lead, rest, h=0.75, color=ORG, fill=ORG_BG):
    box(s, L, y, CW, h, fill, color, 1)
    rule(s, L, y + 0.12, 0.06, color, h - 0.24)
    text(s, L + 0.35, y, CW - 0.6, h, ((lead, {'bold': True, 'color': color}), (rest, {})), size=13.5, anchor=MIDDLE)

def arrow(s, x, y, color=RED, size=20): text(s, x, y, 0.4, 0.5, '→', size=size, color=color, bold=True, align=CENTER)

old = list(prs.slides)

# ================= 1. Title =================
s = old[0]
for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
s.background.fill.solid(); s.background.fill.fore_color.rgb = C(BG)
oval(s, 8.05, 0.45, 5.2, RED_BG)
pill(s, L, 0.85, 'SOCIAL & SUSTAINABILITY PROJECT', size=11)
text(s, L, 1.6, 8.0, 2.2, [(('Digital Loan App', {}),), (('Harassment', {'color': RED}), (' in India', {}))], size=44, bold=True, spacing=0.95)
rule(s, 0.95, 3.35, 1.3)
text(s, L, 3.6, 7.6, 0.9, ['Framing the crisis through the UN SDGs,', 'and our solution: LoanShield.'], size=17, color=SOFT)
for i, n in enumerate([1, 8, 10, 16]): tile(s, 0.95 + i * 0.68, 4.95, n, SDG[n], 0.55, 15)
text(s, 3.75, 4.95, 3.5, 0.55, 'Four SDGs, one problem', size=12, color=MUTE, anchor=MIDDLE)
team = [('Karthik J', '23BRS1191'), ('K Annu Prasath', '24BCE5409'), ('Pasunuru Sai Vivek', '24BCE1239'), ('Srinidhi Vishal Chejarla', '24BCE1493')]
for i, (nm, reg) in enumerate(team):
    x = 0.95 + i * 3.05
    rule(s, x, 6.1, 2.85, EDGE, 0.02)
    text(s, x, 6.22, 2.85, 0.35, nm, size=12.5, bold=True)
    text(s, x, 6.55, 2.85, 0.3, reg, size=10.5, color=ORG)
box(s, 9.05, 0.75, 3.15, 5.0, '1A1414', EDGE, 2, 0.1)
box(s, 9.2, 0.9, 2.85, 4.7, '201818', None, adj=0.08)
box(s, 10.2, 1.0, 0.85, 0.1, EDGE, None, adj=0.5)
text(s, 9.4, 1.25, 2.5, 0.3, 'Recovery Agent', size=11, bold=True)
text(s, 9.4, 1.5, 2.5, 0.25, 'Today, 10:42 pm', size=8.5, color=MUTE)
for i, m in enumerate(['Pay ₹6,200 in 2 hours or your photo goes to all your contacts.', 'FIR ready. Police will arrest you tomorrow.', 'Your family will know you are a fraud.']):
    y = 1.95 + i * 1.05
    box(s, 9.4, y, 2.45, 0.85, '2E1A1A', RED_DK, 1, 0.18)
    text(s, 9.55, y + 0.08, 2.2, 0.72, m, size=10, color='EAD9DA', anchor=MIDDLE)
oval(s, 11.6, 1.75, 0.42, RED); text(s, 11.6, 1.75, 0.42, 0.42, '3', size=12, bold=True, align=CENTER, anchor=MIDDLE)
notes(s, 'Introduce the team and the problem: illegal loan apps harass borrowers with threats like these. Part 1 frames it through four SDGs; Part 2 presents our solution, LoanShield.')
page[0] = 1

# ================= 2. The crisis =================
s = frame('THE CRISIS', 'Illegal digital lending traps borrowers', 'tr', fresh=old[1])
flow = [('Quick loans attract the vulnerable', 'Instant approval, no paperwork, aimed at people banks turn away.'),
        ('Hidden fees deepen debt', 'Large cuts before the money arrives, then short deadlines and penalties.'),
        ('Aggressive recovery tactics', 'Threats, blackmail and shaming when the borrower cannot pay.')]
w = (CW - 2 * 0.55) / 3
for i, (h, b) in enumerate(flow):
    x = L + i * (w + 0.55)
    card(s, x, 2.3, w, 2.2, h, b, hc=INK, num=i + 1, hs=15, bs=11.5)
    if i < 2: arrow(s, x + w + 0.07, 3.1)
banner(s, 5.0, 'The harm spreads: ', 'families, workplaces and whole communities lose trust in digital finance.', 0.85, RED, RED_BG)

# ================= 3. Why it matters =================
s = frame('WHY IT MATTERS', 'Harassment tactics & real-world harm', 'br', fresh=old[2])
harms = [('Threatening calls & messages', 'Abusive language and fake legal notices create constant fear.'),
         ('Contact-list blackmail', 'Personal data is weaponised to shame and pressure borrowers.'),
         ('Fear, stress & family conflict', 'Repeated harassment causes psychological harm and tears families apart.'),
         ('Privacy violations', 'Data misuse exposes victims and destroys trust in digital systems.')]
w = (CW - 0.3) / 2
for i, (h, b) in enumerate(harms):
    x, y = L + (i % 2) * (w + 0.3), 2.25 + (i // 2) * 1.45
    card(s, x, y, w, 1.25, h, b, hc=RED if i % 2 == 0 else ORG, hs=14.5, bs=11.5)
x = L
for t in ['Access without protection', 'No institutional accountability', 'Unsustainable digital finance']:
    x += pill(s, x, 5.35, t.upper(), ORG, ORG_BG, 10) + 0.25

# ================= 4. SDG 1 =================
s = frame(None, None, 'tr', fresh=old[3])
tile(s, L, 0.6, 1, SDG[1], 0.95, 30)
text(s, L + 1.2, 0.6, 6, 0.95, [(('SDG 1', {'color': MUTE, 'size': 11}),), (('No Poverty', {'bold': True, 'size': 18}),)], anchor=MIDDLE, gap=0)
text(s, L, 1.85, CW, 0.9, 'Predatory lending creates a poverty trap', size=30, bold=True)
steps = [('High charges', 'Push vulnerable households deeper into debt.'), ('Debt cycles', 'Forced to re-borrow just to survive.'),
         ('Essential cuts', 'Less spending on food, health and education.'), ('Financial ruin', 'Poverty becomes impossible to escape.')]
w = (CW - 3 * 0.3) / 4
for i, (h, b) in enumerate(steps):
    card(s, L + i * (w + 0.3), 3.0, w, 1.7, h, b, hc=SDG[1] if i == 3 else INK, num=i + 1, ncol=SDG[1], hs=14, bs=11)
banner(s, 5.15, 'SDG 1: ', 'predatory lending does not just cause debt. It makes poverty structurally harder to escape.', 0.85, SDG[1], RED_BG)

# ================= 5. SDG 8 & 10 =================
s = frame(None, None, 'bl', fresh=old[4])
text(s, L, 0.6, CW, 0.9, 'Unsafe lending undermines fair economic participation', size=28, bold=True)
w = (CW - 0.4) / 2
for i, (n, name, head, body) in enumerate([(8, 'Decent Work & Economic Growth', 'Livelihoods at risk', 'Harassment calls at work disrupt employment. Constant stress reduces productivity. Predatory debt makes decent work meaningless.'),
                                           (10, 'Reduced Inequalities', 'Inequality widens', 'Vulnerable and marginalised groups face the highest risk of exploitation. Without safe credit, existing inequalities only grow.')]):
    x = L + i * (w + 0.4)
    box(s, x, 1.9, w, 4.6, CARD, SDG[n], 1.5)
    tile(s, x + 0.35, 2.25, n, SDG[n], 1.0, 30)
    text(s, x + 1.6, 2.25, w - 1.9, 1.0, [(('SDG ' + str(n), {'color': MUTE, 'size': 11}),), ((name, {'bold': True, 'size': 15}),)], anchor=MIDDLE, gap=0)
    text(s, x + 0.35, 3.65, w - 0.7, 0.5, head, size=20, bold=True, color=SDG[n] if n == 10 else 'E2557A')
    text(s, x + 0.35, 4.3, w - 0.7, 2.0, body, size=13, color=SOFT)

# ================= 6. SDG 16 =================
s = frame(None, None, 'tr', fresh=old[5])
tile(s, L, 0.6, 16, SDG[16], 0.95, 28)
text(s, L + 1.2, 0.6, 7, 0.95, [(('SDG 16', {'color': MUTE, 'size': 11}),), (('Peace, Justice & Strong Institutions', {'bold': True, 'size': 18}),)], anchor=MIDDLE, gap=0)
text(s, L, 1.85, CW, 0.9, 'Threats, data abuse & eroding trust', size=30, bold=True)
rows = [('Threats & fake legal notices', 'Create constant fear and insecurity, undermining personal safety and dignity.'),
        ('Privacy violations', 'Misuse of personal data weakens trust in digital institutions and financial systems.'),
        ('No accessible justice', 'Victims lack clear pathways for accountability, leaving exploiters unchallenged.')]
w = (CW - 2 * 0.3) / 3
for i, (h, b) in enumerate(rows):
    card(s, L + i * (w + 0.3), 3.0, w, 1.8, h, b, hc='5FB0E8', hs=14.5, bs=11.5)
banner(s, 5.25, 'Without strong institutions ', 'to regulate lending and protect data, digital finance becomes a tool of coercion.', 0.85, '5FB0E8', '0C1C28')

# ================= 7. One problem, four SDGs =================
s = frame('ONE PROBLEM, FOUR SDGs', None, 'br', fresh=old[6])
text(s, L, 1.05, CW, 1.6, [(('Digital finance must', {}),), (('empower', {'color': RED}), (', not exploit.', {}))], size=36, bold=True, spacing=0.95, gap=0)
items = [(1, 'Deepens poverty & debt traps'), (8, 'Disrupts livelihoods & work'), (10, 'Widens existing inequalities'), (16, 'Erodes trust & justice')]
w = (CW - 3 * 0.3) / 4
for i, (n, t) in enumerate(items):
    x = L + i * (w + 0.3)
    box(s, x, 3.0, w, 2.0, CARD, EDGE)
    rule(s, x, 3.0, w, SDG[n], 0.08)
    tile(s, x + 0.3, 3.35, n, SDG[n], 0.7, 20)
    text(s, x + 0.3, 4.2, w - 0.6, 0.7, t, size=13.5, bold=True)
text(s, L, 5.4, CW, 1.0, 'Loan-app harassment is not just a fintech issue. It is a sustainable development crisis. Safe, fair and accountable digital finance is essential for inclusive growth in India.', size=13, color=SOFT)

# ================= 8. Problem recap =================
s = frame('PROBLEM RECAP', 'Where the problem leaves victims', 'bl')
steps = [('Easy loan, hidden cost', '₹5,000 promised, ₹3,500 paid, full repayment in 7 days.'),
         ('Threats & blackmail', 'Contact lists, morphed photos and fake police notices.'),
         ('Silence & rebrand', 'The victim pays and stays quiet. The app returns under a new name.')]
w = (CW - 2 * 0.55) / 3
for i, (h, b) in enumerate(steps):
    x = L + i * (w + 0.55)
    card(s, x, 2.25, w, 1.85, h, b, hc=INK, edge=RED_DK, num=i + 1, hs=15, bs=11.5)
    if i < 2: arrow(s, x + w + 0.07, 2.95)
banner(s, 4.5, 'What is missing: ', 'a way for victims to fight back, together, with proof.', 0.85)
for i, n in enumerate([1, 8, 10, 16]): tile(s, L + i * 0.72, 5.75, n, SDG[n], 0.58, 15)
text(s, L + 3.05, 5.75, 8, 0.58, 'Our solution answers each stage of it.', size=12.5, color=MUTE, anchor=MIDDLE)
notes(s, 'Bridge from Part 1. One line per box: how people get trapped, how they are threatened, why the operators escape. Then: what is missing is a way for victims to fight back with proof.')

# ================= 9. Objectives =================
s = frame('OBJECTIVES', 'What our solution sets out to do', 'tr')
objs = [('Calm the victim', 'Answer every threat with the facts and the law it breaks.'),
        ('Build evidence', 'Seal each message as tamper-proof proof.'),
        ('Protect the contacts', 'Warn family and friends before the scammers reach them.'),
        ('Unmask the operator', 'Link rebranded apps by UPI ID, phone number and wording.'),
        ('Offer a fair alternative', 'Small, fair loans through SHGs and microfinance groups.'),
        ('Reach everyone', 'Tamil, Hindi and English; help for people who can’t use phones.')]
w = (CW - 2 * 0.3) / 3
for i, (h, b) in enumerate(objs):
    card(s, L + (i % 3) * (w + 0.3), 2.2 + (i // 3) * 2.15, w, 1.9, h, b, hc=INK, num=i + 1, ncol=RED if i < 4 else ORG, hs=14.5, bs=11.5)
notes(s, 'Six objectives. The first four are what the website does; five and six are delivered through community partners.')

# ================= 10. Our solution =================
s = frame('OUR SOLUTION', None, 'r')
text(s, L, 1.05, 5.8, 0.8, 'LoanShield', size=40, bold=True)
text(s, L, 1.9, 5.8, 0.7, (('Their threats become ', {}), ('evidence.', {'color': RED})), size=24, bold=True)
rule(s, L + 0.05, 2.75, 1.0)
text(s, L, 2.95, 5.6, 1.1, 'A community system with a website at its centre. Other tools try to block the threats. LoanShield turns each one into a fact for the victim and evidence against the operator.', size=12.5, color=SOFT)
for i, (h, b) in enumerate([('Flip', 'Every threat is answered with what is actually true.'), ('Seal', 'Every message becomes tamper-proof evidence.'), ('Link', 'Every UPI ID leads back to one operator.')]):
    y = 4.3 + i * 0.75
    box(s, L, y, 1.0, 0.55, RED, None, adj=0.5)
    text(s, L, y, 1.0, 0.55, h, size=13, bold=True, align=CENTER, anchor=MIDDLE)
    text(s, L + 1.2, y, 4.6, 0.55, b, size=12.5, anchor=MIDDLE)
picture(s, IMG + 'n_hero_c.jpg', 7.05, 1.3, w=5.3)
text(s, 7.05, 4.6, 5.3, 0.4, (('Live  ', {'bold': True, 'color': ORG}), (LINK, {'color': SOFT})), size=11.5)
notes(s, 'The one-line idea: their threats become evidence. Three verbs: flip, seal, link. Show the live website if time allows.')

# ================= 11. How it works =================
s = frame('HOW IT WORKS', 'Five stages, one loop', 'br')
st = [('Prevent', 'Community', ['Real-cost camps at SHGs and colleges', 'Fair small loans via SHGs']),
      ('Detect', 'Website', ['Links reports by UPI ID and number', 'Public only after 3 people report']),
      ('Protect', 'Website + NGO', ['Flips each threat into fact and law', 'Warns contacts in 3 languages']),
      ('Respond', 'Community', ['Legal-aid clinics help file', 'Leaked-photo takedowns by NGOs']),
      ('Reform', 'Community', ['Public dashboard of operators', 'Reports to app stores']) ]
W = (CW - 4 * 0.28) / 5
for i, (n, who, items) in enumerate(st):
    x = L + i * (W + 0.28); col = RED if 'Website' in who else ORG
    box(s, x, 2.15, W, 3.25, CARD, col)
    dot(s, x + 0.22, 2.37, i + 1, col, 0.46, 12)
    text(s, x + 0.8, 2.37, W - 0.9, 0.46, n, size=15, bold=True, anchor=MIDDLE)
    text(s, x + 0.22, 3.02, W - 0.4, 0.3, who.upper(), size=9, color=col, bold=True)
    text(s, x + 0.22, 3.4, W - 0.4, 1.9, ['• ' + t for t in items], size=10.5, color=SOFT, gap=6)
    if i < 4: arrow(s, x + W - 0.06, 3.5, MUTE, 15)
banner(s, 5.7, '↺  Loop: ', 'hotspots found in Reform decide where the next Prevent camps run.', 0.7)
text(s, L, 6.55, CW, 0.3, (('●  ', {'color': RED}), ('Carried by the website        ', {}), ('●  ', {'color': ORG}), ('Carried by people and partners', {})), size=10, color=MUTE)
notes(s, 'Walk left to right. The website handles Detect and Protect. Prevent, Respond and Reform are run by SHGs, NGO counsellors, legal-aid clinics and student volunteers.')

# ================= 12. The website =================
s = frame('THE WEBSITE: ONE PART OF THE SYSTEM', 'Paste a threat. Watch it flip.', 'tr')
calls = [('Speak, paste or screenshot', 'Say what happened; the key details are pulled out.'),
         ('Fear → Fact', 'Each threat card turns over to show the truth and the law.'),
         ('Protect me in one tap', 'Seals it, links the operator, writes the warning and complaint.'),
         ('Case tracker', 'Evidence, laws, operator, warning and complaint in one view.')]
for i, (h, b) in enumerate(calls):
    card(s, L, 2.15 + i * 1.18, 3.75, 1.05, h, b, hc=RED if i == 1 else INK, hs=12.5, bs=10)
w, h = picture(s, IMG + 'n_cards.jpg', 5.05, 2.2, w=7.3)
picture(s, IMG + 'n_case.jpg', 5.05, 2.2 + h + 0.3, w=7.3)
notes(s, 'Demo: paste the example message, show the cards flipping, then the case tracker filling up. Everything is checked on the phone; nothing is uploaded unless the victim shares it.')

# ================= 13. Tech stack =================
s = frame('TECH STACK', 'Simple tools, chosen to reach everyone', 'bl')
stack = [('Website', 'Opens in any phone browser. No app to install.', 'HTML · CSS · JavaScript'),
         ('Reading threats', 'Spots threat words in English, Hindi and Tamil.', 'Keyword detection + AI for screenshots'),
         ('Voice', 'Speak instead of typing; hear the facts read out.', 'Phone’s built-in speech tools'),
         ('Evidence seal', 'A digital fingerprint shows if anything changed.', 'SHA-256 hashing'),
         ('Linking operators', 'Matches UPI IDs and numbers across reports.', 'Pattern matching'),
         ('Privacy & hosting', 'Checked on the phone; shared only by choice.', 'Free hosting on GitHub Pages')]
w = (CW - 2 * 0.3) / 3
for i, (h, b, t) in enumerate(stack):
    x, y = L + (i % 3) * (w + 0.3), 2.2 + (i // 3) * 2.15
    box(s, x, y, w, 1.9)
    text(s, x + 0.25, y + 0.24, w - 0.5, 0.4, h, size=14.5, color=RED, bold=True)
    text(s, x + 0.25, y + 0.72, w - 0.5, 0.7, b, size=11.5)
    rule(s, x + 0.25, y + 1.38, w - 0.5, LINE, 0.01)
    text(s, x + 0.25, y + 1.48, w - 0.5, 0.3, t, size=9.5, color=MUTE)
notes(s, 'Keep it light: everything runs in a normal phone browser, needs no app and costs almost nothing to run. The point is reach, not complexity.')

# ================= 14. Two ways in =================
s = frame('FOR PEOPLE WHO STRUGGLE WITH TECH', 'Two ways in: do it yourself, or get assistance', 'br')
acc = [('Do it yourself', 'Paste, speak or upload a screenshot. About a minute.', RED),
       ('Someone is with me', 'Helper mode: big text, 4 simple steps, read aloud.', ORG),
       ('Call me back', 'A trained NGO volunteer calls in their language.', ORG),
       ('Their language, out loud', 'Tamil, Hindi or English; facts are spoken.', ORG)]
for i, (h, b, c) in enumerate(acc):
    x, y = L + (i % 2) * 3.05, 2.2 + (i // 2) * 1.75
    card(s, x, y, 2.85, 1.55, h, b, hc=c, hs=13.5, bs=10.5)
banner(s, 5.85, 'No smartphone at all? ', 'Student volunteers run weekly help desks at SHG meetings and do every step for them.', 0.75)
picture(s, IMG + 'n_start.jpg', 7.25, 2.25, w=5.1)
notes(s, 'The first people trapped are those least used to smartphones. The website asks first: do it yourself or get assistance. Help desks at SHG meetings reach people with no smartphone at all.')

# ================= 15. Leaked photo =================
s = frame('WHEN A MORPHED PHOTO LEAKS', 'Evidence is not enough. Dignity comes first.', 'tr')
cols = [('Before it spreads', ORG, ['Contact Shield warns family first: “any photo of me is fake”', 'The photo arrives already discredited']),
        ('Within hours', RED, ['Ready takedown request for each platform', 'Platforms must act within 2 hours of a complaint (IT Rules, 2026)', 'StopNCII blocks re-uploads; the image never leaves the phone']),
        ('For the person', '5FB0E8', ['An NGO counsellor does the takedowns', 'A woman volunteer on request', 'They never have to see it again'])]
w = (CW - 2 * 0.3) / 3
for i, (h, c, items) in enumerate(cols):
    x = L + i * (w + 0.3)
    box(s, x, 2.15, w, 2.95, CARD, EDGE); rule(s, x, 2.15, w, c, 0.08)
    text(s, x + 0.25, 2.45, w - 0.5, 0.45, h, size=15, color=c, bold=True)
    text(s, x + 0.25, 3.0, w - 0.5, 2.0, ['• ' + t for t in items], size=11, color=SOFT, gap=7)
banner(s, 5.45, 'On the website: ', '“This photo is fake. Making it is the crime, not being in it.” One button opens the whole checklist.', 0.85, RED, RED_BG)
notes(s, 'Answer to "isn’t the leaked photo still humiliating?": yes, so the response is takedown plus support, not just evidence. Warn contacts first, take it down within hours, block re-uploads, and let a counsellor do it so the victim never sees it again. WhatsApp copies cannot be deleted from other phones, which is why the early warning matters most.')

# ================= 16. Novelty =================
s = frame('NOVELTY', 'What is new in our approach', 'bl')
text(s, L, 2.05, 4.9, 0.3, 'USUAL APPROACH', size=10, color=MUTE, bold=True)
text(s, 6.55, 2.05, 5.8, 0.3, 'LOANSHIELD', size=10, color=RED, bold=True)
rows = [('Spam blockers hide the message. The victim is still scared.', 'Flips fear into facts', 'Each threat is answered with the truth and the law it breaks.'),
        ('More threats mean more fear. Each victim complains alone.', 'Every threat builds the case', 'Messages are sealed as evidence; contacts are warned first.'),
        ('One app is reported. It returns next week with a new name.', 'Follows the money', 'Reports sharing a UPI ID or number join into one operator.')]
for i, (o, h, b) in enumerate(rows):
    y = 2.45 + i * 1.28
    box(s, L, y, 5.0, 1.08, CARD, EDGE)
    text(s, L + 0.25, y, 4.5, 1.08, o, size=11.5, color=MUTE, anchor=MIDDLE)
    arrow(s, 5.98, y + 0.28, RED, 20)
    box(s, 6.55, y, R - 6.55, 1.08, RED_BG, RED)
    text(s, 6.8, y + 0.16, R - 7.1, 0.4, h, size=14.5, bold=True)
    text(s, 6.8, y + 0.56, R - 7.1, 0.45, b, size=11, color=SOFT)
text(s, L, 6.4, CW, 0.45, 'Other tools defend one phone. LoanShield turns every victim’s report into one case against the operator.', size=13, color=ORG, bold=True)
notes(s, 'If asked "isn’t this a spam checker?": spam checkers block a number. We answer the fear, build evidence and link victims together by the money trail.')

# ================= 17. SDG impact =================
s = frame('SDG IMPACT', 'How the solution serves the SDGs', 'tr')
box(s, L, 2.15, 5.7, 4.45, '2A0D1C', SDG[10], 1.5)
tile(s, L + 0.35, 2.5, 10, SDG[10], 1.0, 28)
text(s, L + 1.6, 2.5, 3.9, 1.0, [(('MAJOR SDG  ·  TARGET 10.2', {'color': 'F06AA6', 'size': 10, 'bold': True}),), (('Reduced Inequalities', {'bold': True, 'size': 17}),)], anchor=MIDDLE, gap=2)
text(s, L + 0.35, 3.75, 5.0, 0.6, 'These apps exploit the people left behind, financially and digitally. LoanShield serves them first.', size=12, color=INK, bold=True)
text(s, L + 0.35, 4.5, 5.0, 2.0, ['•  Assisted help: a helper, a call-back or an SHG help desk', '•  Their language, read aloud: Tamil, Hindi, English', '•  No app, no sign-up: works on any cheap phone'], size=11.5, color='E8CCD9', gap=8)
for i, (n, h, b) in enumerate([(1, 'No Poverty', 'Fair SHG loans and real-cost checks break the debt trap.'), (16, 'Peace & Justice', 'Sealed evidence and legal-aid help give victims access to justice.'), (8, 'Decent Work', 'Contact warnings keep threats away from employers.')]):
    x, y = 6.9, 2.15 + i * 1.55
    box(s, x, y, R - x, 1.35)
    tile(s, x + 0.25, y + 0.32, n, SDG[n], 0.7, 20)
    text(s, x + 1.2, y + 0.22, R - x - 1.4, 0.4, 'SDG ' + str(n) + '  ·  ' + h, size=13.5, bold=True)
    text(s, x + 1.2, y + 0.62, R - x - 1.4, 0.6, b, size=10.5, color=SOFT)
notes(s, 'Lead with SDG 10, target 10.2: empower and include everyone regardless of income or ability. These apps target the poorest and least digitally literate; LoanShield is designed for them first (assisted help, local languages, read aloud, no app). SDG 1, 16 and 8 are supporting goals. Line: "These apps exploit exactly the people who are left behind, financially and digitally. LoanShield is designed for them first, which is why SDG 10 is our main goal."')

# ================= 18. Funding & partners =================
s = frame('FUNDING & PARTNERS', 'Run with cyber-safety NGOs', 'br')
w = (CW - 2 * 0.3) / 3
for i, (h, b) in enumerate([('CyberPeace Foundation', 'Cyber-safety NGO. Pilot funding, counsellor network and victim helpdesk.'), ('Responsible Netism', 'Online-safety NGO. Awareness workshops in colleges and communities.'), ('Cyber Society of India', 'Chennai-based cyber awareness body. Volunteer training for our pilot.')]):
    card(s, L + i * (w + 0.3), 2.15, w, 1.95, h, b, hc=ORG, hs=14.5, bs=11)
w2 = (CW - 0.3) / 2
for i, (h, b) in enumerate([('CSR grants', 'from responsible fintech companies and banks'), ('Law-college clinics', 'pro-bono legal help for filing'), ('SHGs & microfinance groups', 'the fair-loan alternative'), ('Student volunteers', 'NSS and social clubs run camps')]):
    x, y = L + (i % 2) * (w2 + 0.3), 4.4 + (i // 2) * 0.88
    box(s, x, y, w2, 0.7)
    text(s, x + 0.25, y, w2 - 0.5, 0.7, ((h + ':  ', {'bold': True}), (b, {'color': SOFT})), size=11.5, anchor=MIDDLE)
text(s, L, 6.3, CW, 0.3, 'Partners we plan to approach for the pilot. Running cost is low: the website is free to host.', size=10.5, color=MUTE)
notes(s, 'Funding and people come from cyber-safety NGOs and CSR, not from government programmes. These are organisations we plan to approach; none has committed yet.')

# ================= 19. Challenges & limitations =================
s = frame('CHALLENGES & LIMITATIONS', 'What makes it hard, and what it can’t do yet', 'tr')
w = (CW - 0.4) / 2
text(s, L, 2.05, w, 0.3, 'CHALLENGES  →  HOW WE HANDLE THEM', size=10, color=ORG, bold=True)
for i, (a, b) in enumerate([('Apps rebrand in weeks', 'link by UPI ID and number, not app name'), ('Shame stops reporting', 'private by default, NGO counsellors'), ('False or spiteful reports', 'public only after 3 different people'), ('Low digital literacy', 'assisted help, voice and local languages')]):
    y = 2.45 + i * 0.98
    box(s, L, y, w, 0.84)
    text(s, L + 0.25, y, w - 0.5, 0.84, [((a, {'bold': True}),), (('→  ' + b, {'color': SOFT, 'size': 10.5}),)], size=12, anchor=MIDDLE, gap=2)
x2 = L + w + 0.4
text(s, x2, 2.05, w, 0.3, 'LIMITATIONS', size=10, color=RED, bold=True)
for i, t in enumerate(['Keyword scanner can miss new wording', 'Tested on 50 team-written messages, not field data', 'Evidence is stored only on the victim’s own device', 'Legal mapping is a guide, not legal advice', 'People-side stages depend on NGO partners']):
    y = 2.45 + i * 0.78
    box(s, x2, y, w, 0.64, '1F1216', RED_DK)
    text(s, x2 + 0.25, y, w - 0.5, 0.64, t, size=11.5, color='E0CFD1', anchor=MIDDLE)
notes(s, 'Be upfront about limits. Scanner results on our test set: 100% recall, 86% precision, but on team-written messages.')

# ================= 20. Future scope =================
s = frame('FUTURE SCOPE', 'Where LoanShield goes next', 'bl')
fut = [('Android app', 'Share any message straight in; log late-night calls automatically.'), ('Smarter reading', 'A trained model for Tamil, Telugu and Bengali threats.'),
       ('WhatsApp bot', 'Forward a threat to one number; get the facts back as a voice note.'), ('Voice notes', 'Turn threatening voice messages into text evidence.'),
       ('90-day pilot', 'One SHG federation and one cyber-safety NGO in Chennai.'), ('Open data for NGOs', 'Share operator links with partners to speed up takedowns.')]
w = (CW - 2 * 0.3) / 3
for i, (h, b) in enumerate(fut):
    card(s, L + (i % 3) * (w + 0.3), 2.2 + (i // 3) * 2.1, w, 1.85, h, b, hc=ORG, edge=ORG if i == 4 else EDGE, num=i + 1, ncol=ORG if i == 4 else RED_DK, hs=14.5, bs=11)
notes(s, 'The pilot is the most important next step: it gives real data to replace our test set.')

# ================= 21. Thank you =================
s = frame(None, None, 'r')
oval(s, 8.4, 0.9, 4.6, RED_BG)
pill(s, L, 1.0, 'THANK YOU', size=11)
text(s, L, 1.65, 8, 1.9, [(('Safer loans.', {}),), (('Stronger people.', {'color': RED}),)], size=46, bold=True, spacing=0.95, gap=0)
rule(s, 0.95, 3.65, 1.3)
for i, (nm, reg) in enumerate(team):
    x, y = L + (i % 2) * 3.4, 4.1 + (i // 2) * 0.85
    text(s, x, y, 3.3, 0.35, nm, size=13, bold=True)
    text(s, x, y + 0.36, 3.3, 0.3, reg, size=10.5, color=ORG)
box(s, 9.0, 2.2, 3.4, 2.6, CARD, EDGE, 1, 0.08)
oval(s, 9.3, 2.5, 0.7, RED); text(s, 9.3, 2.5, 0.7, 0.7, '→', size=20, bold=True, align=CENTER, anchor=MIDDLE)
text(s, 9.3, 3.4, 2.9, 0.4, 'Try LoanShield', size=15, bold=True)
text(s, 9.3, 3.85, 2.9, 0.8, LINK, size=10.5, color=SOFT)
notes(s, 'Close and invite questions.')

prs.save(OUT); print('saved', len(prs.slides), 'slides')
