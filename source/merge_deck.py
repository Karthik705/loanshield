# Continues the team's Part 1 deck in the same template and appends the solution slides.
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

SRC, OUT = sys.argv[1], sys.argv[2]
IMG = '/home/claude/loanshield/deckimg/'
LINK = 'https://karthik705.github.io/loanshield/'
C = lambda h: RGBColor.from_string(h)
BG_WARM, BG_INK, BG_BLUE = '120A08', '0B0B0F', '080D12'
INK, MUTE, CARD, EDGE = 'F1FAEE', '9A9AAA', '1A1A24', '2E2E3A'
RED, ORG = 'E63946', 'F4A261'
SDG = {1: 'E5243B', 8: 'A21942', 10: 'DD1367', 16: '00689D'}
F = 'Calibri'

prs = Presentation(SRC)
BLANK = prs.slide_layouts[6]

def new(bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = C(bg)
    return s

def text(s, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = p if isinstance(p, tuple) and isinstance(p[0], tuple) else (p if isinstance(p, tuple) else ((p, {}),))
        if isinstance(p, str): segs = ((p, {}),)
        for t, o in segs:
            r = para.add_run(); r.text = t; f = r.font
            f.name = F; f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold); f.color.rgb = C(o.get('color', color))
        if i: para.space_before = Pt(6)
    return tb

def eyebrow(s, t): text(s, 1.0, 0.5, 8, 0.35, t, size=13, color=RED, bold=True)
def title(s, t, y=0.95, size=40): text(s, 1.0, y, 11.3, 1.3, t, size=size, bold=True)

def box(s, x, y, w, h, fill=CARD, edge=EDGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.5, adj=0.12):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    if edge: sh.line.color.rgb = C(edge); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = adj
    sh.text_frame.text = ''
    return sh

def card(s, x, y, w, h, head, body, hc=RED, edge=EDGE, fill=CARD, hs=18, bs=13.5, align=PP_ALIGN.LEFT):
    box(s, x, y, w, h, fill, edge)
    text(s, x + 0.25, y + 0.22, w - 0.5, 0.5, head, size=hs, color=hc, bold=True, align=align)
    text(s, x + 0.25, y + 0.24 + hs / 72 * 1.35, w - 0.5, h - 0.7, body, size=bs, color=MUTE, align=align)

def picture(s, path, x, y, w=None, h=None, edge=EDGE):
    iw, ih = Image.open(path).size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    box(s, x - 0.06, y - 0.06, w + 0.12, h + 0.12, '1A1A24', edge, adj=0.04)
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h

def notes(s, t): s.notes_slide.notes_text_frame.text = t

def num_tile(s, x, y, n, color, size=0.7, fs=26):
    b = box(s, x, y, size, size, color, None, adj=0.15)
    text(s, x, y, size, size, str(n), size=fs, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def arrow(s, x, y, color=RED, size=22): text(s, x, y, 0.4, 0.5, '→', size=size, color=color, bold=True, align=PP_ALIGN.CENTER)

# ---------------- 8. Problem recap ----------------
s = new(BG_WARM); eyebrow(s, 'PROBLEM RECAP')
title(s, ['Where the problem leaves victims'])
steps = [('Easy loan, hidden cost', '₹5,000 promised, ₹3,500 paid, full repayment in 7 days.'),
         ('Threats & blackmail', 'Contact lists, morphed photos and fake police notices.'),
         ('Silence & rebrand', 'The victim pays and stays quiet. The app returns under a new name.')]
for i, (h, b) in enumerate(steps):
    x = 1.0 + i * 3.85
    card(s, x, 2.25, 3.35, 1.75, h, b, hc=RED, edge=RED)
    if i < 2: arrow(s, x + 3.38, 2.85)
box(s, 1.0, 4.35, 11.05, 0.9, '1A1A24', ORG)
text(s, 1.3, 4.35, 10.5, 0.9, (('What is missing: ', {'bold': True, 'color': ORG}), ('a way for victims to fight back, together, with proof.', {})), size=18, anchor=MSO_ANCHOR.MIDDLE)
for i, n in enumerate([1, 8, 10, 16]):
    num_tile(s, 1.0 + i * 0.85, 5.75, n, SDG[n], 0.62, 20)
text(s, 4.55, 5.75, 7.5, 0.62, 'One problem, four SDGs. Our solution answers each stage of it.', size=15, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, 'Bridge from Part 1. One line per box: how people get trapped, how they are threatened, why the operators escape. Then: what is missing is a way for victims to fight back with proof.')

# ---------------- 9. Objectives ----------------
s = new(BG_INK); eyebrow(s, 'OBJECTIVES'); title(s, 'What our solution sets out to do')
objs = [('Calm the victim', 'Answer every threat with the facts and the law it breaks.'),
        ('Build evidence automatically', 'Seal each message as tamper-proof proof.'),
        ('Protect the contacts', 'Warn family and friends before the scammers reach them.'),
        ('Unmask the operator', 'Link rebranded apps by UPI ID, phone number and wording.'),
        ('Offer a fair alternative', 'Small, fair loans through SHGs and microfinance groups.'),
        ('Reach everyone', 'Tamil, Hindi and English, WhatsApp for basic phones.')]
for i, (h, b) in enumerate(objs):
    x, y = 1.0 + (i % 3) * 3.75, 2.2 + (i // 3) * 2.15
    box(s, x, y, 3.5, 1.9)
    num_tile(s, x + 0.25, y + 0.25, i + 1, RED if i % 2 == 0 else 'C1121F', 0.55, 18)
    text(s, x + 0.95, y + 0.28, 2.4, 0.55, h, size=17, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 1.0, 3.0, 0.8, b, size=13.5, color=MUTE)
notes(s, 'Six objectives. The first four are what the app does; five and six are delivered through community partners.')

# ---------------- 10. Our solution ----------------
s = new(BG_WARM); eyebrow(s, 'OUR SOLUTION')
text(s, 1.0, 0.95, 5.5, 0.8, 'LoanShield', size=40, bold=True)
text(s, 1.0, 1.75, 5.5, 0.9, (('Their threats become ', {}), ('evidence.', {'color': RED})), size=28, bold=True)
text(s, 1.0, 2.85, 5.3, 1.0, 'A community system with a website at its centre. Other tools try to block the threats. LoanShield turns each one into a fact for the victim and evidence against the operator.', size=15, color=MUTE)
pts = [('Flip', 'Every threat is answered with what is actually true.'), ('Seal', 'Every message becomes tamper-proof evidence.'), ('Link', 'Every UPI ID leads back to one operator.')]
for i, (h, b) in enumerate(pts):
    y = 4.05 + i * 0.78
    box(s, 1.0, y, 0.95, 0.58, RED, None, adj=0.25)
    text(s, 1.0, y, 0.95, 0.58, h, size=15, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.15, y, 4.2, 0.58, b, size=15, anchor=MSO_ANCHOR.MIDDLE)
picture(s, IMG + 'n_hero_c.jpg', 6.75, 1.2, w=5.6)
text(s, 6.75, 4.85, 5.6, 0.4, (('Live: ', {'bold': True, 'color': ORG}), (LINK, {'color': INK})), size=13)
notes(s, 'The one-line idea: their threats become evidence. Three verbs: flip, seal, link. Show the live website if time allows.')

# ---------------- 11. How it works ----------------
s = new(BG_INK); eyebrow(s, 'HOW IT WORKS'); title(s, 'Five stages, one loop')
st = [('Prevent', 'Community', ['Real-cost camps at SHGs and colleges', 'Fair small loans via SHGs and microfinance']),
      ('Detect', 'Website', ['Link reports by UPI ID, number, wording', 'Public only after 3 people report']),
      ('Protect', 'Website + NGO', ['Flip each threat into fact and law', 'Warn contacts in 3 languages']),
      ('Respond', 'Community', ['Legal-aid clinics help file', 'One bundled case per operator']),
      ('Reform', 'Community', ['Public dashboard of operators', 'Reports to app stores, payment apps'])]
W = 2.05
for i, (n, who, items) in enumerate(st):
    x = 1.0 + i * (W + 0.25); web = 'Website' in who
    col = RED if web else ORG
    box(s, x, 2.15, W, 3.2, CARD, col)
    num_tile(s, x + 0.2, 2.35, i + 1, col, 0.5, 16)
    text(s, x + 0.82, 2.35, W - 0.9, 0.5, n, size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 3.0, W - 0.4, 0.3, who.upper(), size=11, color=col, bold=True)
    text(s, x + 0.2, 3.4, W - 0.35, 1.9, ['• ' + t for t in items], size=12.5, color=MUTE)
    if i < 4: arrow(s, x + W - 0.07, 3.45, MUTE, 18)
box(s, 1.0, 5.65, 11.25, 0.65, '1A1A24', EDGE)
text(s, 1.3, 5.65, 10.8, 0.65, (('↺  Loop: ', {'bold': True, 'color': ORG}), ('hotspots found in Reform decide where the next Prevent camps run.', {})), size=15, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.0, 6.5, 11, 0.35, (('■ ', {'color': RED}), ('Carried by the website     ', {}), ('■ ', {'color': ORG}), ('Carried by people and partners', {})), size=12, color=MUTE)
notes(s, 'Walk left to right. The website handles Detect and Protect. Prevent, Respond and Reform are run by SHGs, NGO counsellors, legal-aid clinics and student volunteers.')

# ---------------- 12. The website ----------------
s = new(BG_BLUE); eyebrow(s, 'THE WEBSITE: ONE PART OF THE SYSTEM'); title(s, 'Paste a threat. Watch it flip.')
calls = [('Fear → Fact', 'Each threat card turns over to show the truth and the law it breaks.'),
         ('Case tracker', 'Sealed evidence, laws, operator, contact warning, complaint.'),
         ('Operator link', 'Same UPI ID as 3 other app names = one operator.')]
for i, (h, b) in enumerate(calls):
    y = 2.2 + i * 1.35
    card(s, 1.0, y, 3.6, 1.15, h, b, hc='4EA8DE', hs=16, bs=12.5)
w, h = picture(s, IMG + 'n_cards.jpg', 5.0, 2.2, w=7.3)
picture(s, IMG + 'n_case.jpg', 5.0, 2.2 + h + 0.25, w=7.3)
notes(s, 'Demo: paste the example message, show the cards flipping, then the case tracker filling up. Everything is checked on the phone; nothing is uploaded unless the victim shares it.')

# ---------------- 13. Novelty ----------------
s = new(BG_WARM); eyebrow(s, 'NOVELTY'); title(s, 'What is new in our approach')
text(s, 1.0, 2.05, 4.8, 0.35, 'USUAL APPROACH', size=12, color=MUTE, bold=True)
text(s, 6.55, 2.05, 5.6, 0.35, 'LOANSHIELD', size=12, color=RED, bold=True)
rows = [('Spam blockers hide the message. The victim is still scared.', 'Flips fear into facts', 'Each threat is answered with the truth and the law it breaks.'),
        ('More threats mean more fear. Each victim complains alone.', 'Every threat builds the case', 'Messages are sealed as evidence; contacts are warned first.'),
        ('One app is reported. It returns next week with a new name.', 'Follows the money', 'Reports sharing a UPI ID or number join into one operator.')]
for i, (old, h, b) in enumerate(rows):
    y = 2.5 + i * 1.3
    box(s, 1.0, y, 4.9, 1.1, '1A1A24', EDGE)
    text(s, 1.25, y, 4.45, 1.1, old, size=14, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, 5.98, y + 0.27, RED, 24)
    box(s, 6.55, y, 5.7, 1.1, '2A0E10', RED)
    text(s, 6.8, y + 0.15, 5.2, 0.4, h, size=17, bold=True)
    text(s, 6.8, y + 0.55, 5.2, 0.5, b, size=13, color='C9C9D3')
text(s, 1.0, 6.5, 11.3, 0.5, 'Other tools defend one phone. LoanShield turns every victim’s report into one case against the operator.', size=15, color=ORG, bold=True)
notes(s, 'If asked "isn’t this a spam checker?": spam checkers block a number. We answer the fear, build evidence and link victims together by the money trail.')

# ---------------- 14. SDG impact ----------------
s = new(BG_INK); eyebrow(s, 'SDG IMPACT'); title(s, 'How the solution serves the SDGs')
box(s, 1.0, 2.15, 5.6, 4.35, '0A1B2A', SDG[16], lw=2.5)
num_tile(s, 1.3, 2.45, 16, SDG[16], 1.05, 34)
text(s, 2.6, 2.45, 3.8, 1.05, ['MAJOR SDG', 'Peace, Justice & Strong Institutions'], size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.3, 3.8, 5.0, 2.6, ['• Access to justice: sealed evidence and legal-aid help to file',
                             '• Less violence and fear: threats answered with facts',
                             '• Accountability: operators named in public, not just apps',
                             '• Privacy protected: checked on the phone, shared only by consent'], size=14, color='C9D6E2')
minor = [(1, 'No Poverty', 'Fair SHG loans break the debt trap; real-cost checks before borrowing.'),
         (8, 'Decent Work', 'Contact Shield keeps threats from reaching employers.'),
         (10, 'Reduced Inequalities', 'Local languages and WhatsApp reach low-literacy users.')]
for i, (n, h, b) in enumerate(minor):
    y = 2.15 + i * 1.5
    box(s, 6.9, y, 5.35, 1.3)
    num_tile(s, 7.1, y + 0.3, n, SDG[n], 0.7, 22)
    text(s, 8.0, y + 0.2, 4.1, 0.4, 'SDG ' + str(n) + ' · ' + h, size=16, bold=True, color=INK)
    text(s, 8.0, y + 0.6, 4.1, 0.65, b, size=12.5, color=MUTE)
notes(s, 'Lead with SDG 16: the core harm is intimidation and no access to justice, and that is what LoanShield attacks directly. SDGs 1, 8 and 10 are supported.')

# ---------------- 15. Funding & partners ----------------
s = new(BG_BLUE); eyebrow(s, 'FUNDING & PARTNERS'); title(s, 'Run with cyber-safety NGOs')
ngos = [('CyberPeace Foundation', 'Cyber-safety NGO. Pilot funding, counsellor network and victim helpdesk support.'),
        ('Responsible Netism', 'Online-safety NGO. Awareness workshops in colleges and communities.'),
        ('Cyber Society of India', 'Chennai-based cyber awareness body. Volunteer training for our Tamil Nadu pilot.')]
for i, (h, b) in enumerate(ngos):
    card(s, 1.0 + i * 3.8, 2.15, 3.5, 1.95, h, b, hc='4EA8DE', hs=17, bs=13)
others = [('CSR grants', 'from responsible fintech companies and banks'), ('Law-college clinics', 'pro-bono legal help for filing'), ('SHGs & microfinance groups', 'the fair-loan alternative'), ('Student volunteers', 'NSS and social clubs run camps')]
for i, (h, b) in enumerate(others):
    x, y = 1.0 + (i % 2) * 5.7, 4.4 + (i // 2) * 0.9
    box(s, x, y, 5.45, 0.72, '1A1A24', EDGE)
    text(s, x + 0.25, y, 5.0, 0.72, ((h + ': ', {'bold': True, 'color': INK}), (b, {})), size=13.5, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.0, 6.35, 11, 0.35, 'Partners we plan to approach for the pilot. Running cost is low: the website is free to host.', size=12, color=MUTE)
notes(s, 'Funding and people come from cyber-safety NGOs and CSR, not from government programmes. These are organisations we plan to approach; none has committed yet.')

# ---------------- 16. Challenges & limitations ----------------
s = new(BG_INK); eyebrow(s, 'CHALLENGES & LIMITATIONS'); title(s, 'What makes it hard, and what it can’t do yet')
text(s, 1.0, 2.05, 5.5, 0.4, 'CHALLENGES → HOW WE HANDLE THEM', size=12, color=ORG, bold=True)
ch = [('Apps rebrand in weeks', 'link by UPI ID and number, not app name'), ('Shame stops reporting', 'private by default, NGO counsellors'),
      ('False or spiteful reports', 'public only after 3 different people'), ('Low digital literacy', 'WhatsApp, voice and local languages')]
for i, (a, b) in enumerate(ch):
    y = 2.5 + i * 0.95
    box(s, 1.0, y, 5.45, 0.8, '1A1A24', EDGE)
    text(s, 1.25, y, 5.0, 0.8, [((a, {'bold': True, 'color': INK}),), (('→ ' + b, {}),)], size=13, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 6.8, 2.05, 5.5, 0.4, 'LIMITATIONS', size=12, color=RED, bold=True)
lim = ['Keyword scanner can miss new wording and other languages', 'Tested on 50 messages written by our team, not field data', 'Evidence is stored only on the victim’s own device', 'Legal mapping is a guide, not legal advice', 'People-side stages depend on NGO partners']
for i, t in enumerate(lim):
    y = 2.5 + i * 0.76
    box(s, 6.8, y, 5.45, 0.62, '1F1216', '5A2630')
    text(s, 7.05, y, 5.0, 0.62, t, size=13, color='D8C4C8', anchor=MSO_ANCHOR.MIDDLE)
notes(s, 'Be upfront about limits. Scanner results on our test set: 100% recall, 86% precision, but on team-written messages.')

# ---------------- 17. Future scope ----------------
s = new(BG_WARM); eyebrow(s, 'FUTURE SCOPE'); title(s, 'Where LoanShield goes next')
fut = [('Android app', 'Share any message straight into LoanShield; log late-night calls automatically.'),
       ('Smarter reading', 'A trained model for Tamil, Telugu and Bengali threats.'),
       ('WhatsApp bot', 'Flip a threat by forwarding it, on any phone.'),
       ('Voice notes', 'Turn threatening voice messages into text evidence.'),
       ('90-day pilot', 'One SHG federation and one cyber-safety NGO in Chennai.'),
       ('Open data for NGOs', 'Share operator links with partner NGOs to speed up takedowns.')]
for i, (h, b) in enumerate(fut):
    x, y = 1.0 + (i % 3) * 3.75, 2.2 + (i // 3) * 2.05
    card(s, x, y, 3.5, 1.8, h, b, hc=ORG, edge=ORG if i == 4 else EDGE)
notes(s, 'The pilot is the most important next step: it gives real data to replace our test set.')

# ---------------- 18. Thank you ----------------
s = new(BG_INK)
text(s, 1.0, 1.3, 8, 0.4, 'THANK YOU', size=13, color=RED, bold=True)
text(s, 1.0, 1.8, 11, 1.4, ['Safer loans.', (('Stronger people.', {'color': RED}),)], size=48, bold=True)
box(s, 1.0, 3.75, 2.4, 0.05, RED, None, MSO_SHAPE.RECTANGLE)
team = ['KARTHIK J - 23BRS1191', 'K ANNU PRASATH - 24BCE5409', 'PASUNURU SAI VIVEK - 24BCE1239', 'SRINIDHI VISHAL CHEJARLA - 24BCE1493']
text(s, 1.0, 4.1, 6.5, 1.8, team, size=16, color=ORG, bold=True)
box(s, 7.6, 4.1, 4.65, 1.5, '1A1A24', EDGE)
text(s, 7.85, 4.2, 4.2, 1.3, [(('Try LoanShield', {'bold': True, 'color': INK, 'size': 16}),), ((LINK, {'color': '4EA8DE'}),)], size=14, anchor=MSO_ANCHOR.MIDDLE)
notes(s, 'Close and invite questions.')

prs.save(OUT); print('saved', len(prs.slides), 'slides')
